from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="222222")
INTEGER_COLUMNS = {"n", "m", "b", "machine_id", "batch_position", "size", "quantity", "type_id",
                   "num_batches", "num_jobs", "total_batches", "max_batch_size", "split_count"}
TEXT_COLUMNS = {"algorithm", "source_label", "job_ids", "validation_passed", "note", "distribution"}


def _typed(value: str, header: str) -> Any:
    if value == "" or header in TEXT_COLUMNS:
        return value
    try:
        return int(value) if header in INTEGER_COLUMNS else float(value)
    except ValueError:
        return value


def _read_csv(path: Path) -> list[list[Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return []
    headers = rows[0]
    return [headers] + [[_typed(value, headers[index]) for index, value in enumerate(row)] for row in rows[1:]]


def _flatten(prefix: str, value: Any, target: dict[str, Any]) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            _flatten(f"{prefix}.{key}" if prefix else key, child, target)
    elif isinstance(value, list):
        target[prefix] = ", ".join(str(item) for item in value)
    else:
        target[prefix] = value


def _style_sheet(sheet, table_name: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
    for column in sheet.columns:
        letter = column[0].column_letter
        header = str(column[0].value or "")
        width = max((len(str(cell.value or "")) for cell in column), default=8) + 2
        sheet.column_dimensions[letter].width = min(60 if header == "job_ids" else 32, max(10, width))
    if sheet.max_row > 1 and sheet.max_column > 0:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)


def _add_rows(workbook: Workbook, name: str, rows: list[list[Any]], table_name: str) -> None:
    sheet = workbook.create_sheet(name)
    for row in rows:
        sheet.append(row)
    _style_sheet(sheet, table_name)


def write_excel_workbook(output_dir: str | Path, types_csv: str | Path) -> Path:
    output_dir = Path(output_dir)
    distribution = output_dir.name
    workbook = Workbook()
    workbook.remove(workbook.active)

    _add_rows(workbook, "Summary", _read_csv(output_dir / "validation_summary.csv"), "SummaryTable")
    _add_rows(workbook, "Types", _read_csv(Path(types_csv)), "TypesTable")

    reference = json.loads((output_dir / "reference_lower_bound.json").read_text(encoding="utf-8"))
    _add_rows(workbook, "Reference LB", [["metric", "value"], *reference.items()], "ReferenceTable")

    algorithms = ("dp_lpt_in", "lpt_dp_in", "dp_bbs_in")
    machine_rows: list[list[Any]] = []
    batch_rows: list[list[Any]] = []
    detail_records: list[dict[str, Any]] = []
    for algorithm in algorithms:
        machines = _read_csv(output_dir / f"{algorithm}_machines.csv")
        batches = _read_csv(output_dir / f"{algorithm}_batches.csv")
        if not machine_rows:
            machine_rows.append(machines[0])
        machine_rows.extend(machines[1:])
        if not batch_rows:
            batch_rows.append(batches[0])
        batch_rows.extend(batches[1:])
        flattened: dict[str, Any] = {}
        payload = json.loads((output_dir / f"{algorithm}_result.json").read_text(encoding="utf-8"))
        _flatten("", payload, flattened)
        detail_records.append(flattened)

    _add_rows(workbook, "Machines", machine_rows, "MachinesTable")
    _add_rows(workbook, "Batches", batch_rows, "BatchesTable")
    detail_headers = list(dict.fromkeys(key for record in detail_records for key in record))
    detail_rows = [detail_headers] + [[record.get(key, "") for key in detail_headers] for record in detail_records]
    _add_rows(workbook, "Algorithm Details", detail_rows, "DetailsTable")

    path = output_dir / f"{distribution}_results.xlsx"
    workbook.save(path)
    return path
